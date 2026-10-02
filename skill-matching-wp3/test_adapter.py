from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
from mahara_data.schemas.profile import CandidateProfile

from cv_adapter import adapter_profil_vers_contrat
from cv_parser import build_profile_from_cv
from main import app


DATA_DIR = Path(__file__).parent / "data"


def test_example_cv_adapts_to_wp1_contract():
    brut = build_profile_from_cv(str(DATA_DIR / "cv_exemple.docx"), "cand_demo")
    adapte = adapter_profil_vers_contrat(brut, governorate_code="TN-71")

    profil = CandidateProfile(**adapte)

    assert profil.years_experience == 5
    assert {skill.label_raw for skill in profil.skills} >= {"Python", "Django", "SQL"}


def test_simple_cv_extracts_skills_and_experience(tmp_path):
    cv_path = tmp_path / "cv_simple.docx"
    document = Document()
    document.add_paragraph("Nadia Test")
    document.add_paragraph("Competences techniques")
    document.add_paragraph("Python, SQL")
    document.add_paragraph("Experience professionnelle")
    document.add_paragraph("2 ans d'experience")
    document.save(cv_path)

    profil = build_profile_from_cv(str(cv_path), "cand_simple")

    assert profil["hard_skills"] == ["Python", "SQL"]
    assert profil["experience_annees"] == 2


def test_empty_cv_returns_empty_profile(tmp_path):
    cv_path = tmp_path / "cv_vide.docx"
    Document().save(cv_path)

    profil = build_profile_from_cv(str(cv_path), "cand_vide")

    assert profil["nom"] == "Nom inconnu"
    assert profil["hard_skills"] == []
    assert profil["soft_skills"] == []
    assert profil["experience_annees"] == 0


def test_corrupt_cv_upload_documents_current_http_500(tmp_path):
    client = TestClient(app, raise_server_exceptions=False)
    response = client.post(
        "/api/v1/parse-cv",
        params={"candidat_id": str(uuid4())},
        files={"fichier": ("cv_corrompu.docx", b"not a DOCX document")},
    )

    assert response.status_code == 500
    assert response.text == "Internal Server Error"
