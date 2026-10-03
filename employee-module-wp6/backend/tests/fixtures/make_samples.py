"""Writes the two sample CVs used by tests/test_cv_import.py (fictional people).

Run from backend/: python -m tests.fixtures.make_samples
"""

from pathlib import Path

import docx

HERE = Path(__file__).parent

PDF_LINES = [
    "AMIRA BEN SALAH",
    "Soudeuse qualifiée",
    "Sfax - amira.bensalah@example.tn - Tél : +216 22 345 678",
    "",
    "Expérience professionnelle",
    "Soudeuse – Société Métallurgique du Sud, Sfax    03/2019 – aujourd'hui",
    "- Soudure à l'arc de structures métalliques",
    "- Travail en équipe avec les monteurs",
    "Aide-soudeuse chez Atelier Ben Ali    2017 - 2019",
    "",
    "Formation",
    "BTS en construction métallique – ISET Sfax, 2017",
    "Baccalauréat Technique – Lycée Hédi Chaker, 2015",
    "",
    "Compétences",
    "Soudure, Lecture de plans, Travail en équipe, Excel",
    "",
    "Langues",
    "Arabe (langue maternelle), Français (courant), Anglais (notions)",
]


def pdf_text(text: str) -> bytes:
    """A PDF string: WinAnsi (cp1252) bytes with ( ) and \\ escaped."""
    raw = text.encode("cp1252")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def write_pdf(path: Path, lines: list[str]) -> None:
    """A one-page text PDF with Helvetica 11pt, written by hand (no PDF library in the stack)."""
    content = b"BT /F1 11 Tf 14 TL 56 780 Td\n" + b"".join(pdf_text(line) + b" Tj T*\n" for line in lines) + b"ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def write_docx(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Curriculum Vitae")
    document.add_heading("Mohamed Trabelsi", level=1)
    contact = document.add_table(rows=3, cols=2)
    for row, (label, value) in zip(contact.rows, [
        ("Adresse", "Rue de la Liberté, Tunis"),
        ("E-mail", "Mohamed.Trabelsi@Example.tn"),
        ("Téléphone", "98.765.432"),
    ]):
        row.cells[0].text, row.cells[1].text = label, value
    document.add_heading("EXPÉRIENCES", level=2)
    document.add_paragraph("Vendeur, Carrefour La Marsa (janvier 2021 - décembre 2023)")
    document.add_paragraph("• Conseil et vente en magasin, encaissement")
    document.add_paragraph("• Conseil aux clients")
    document.add_paragraph("Caissier – Magasin Général — 2019-2020")
    document.add_heading("FORMATION ET DIPLÔMES", level=2)
    document.add_paragraph("Licence en gestion commerciale, ISG Tunis, 2020")
    document.add_paragraph("Baccalauréat Économie et gestion, 2016")
    document.add_heading("COMPÉTENCES", level=2)
    document.add_paragraph("Vente ; Service client ; Excel ; Photoshop")
    document.add_heading("LANGUES", level=2)
    document.add_paragraph("Arabe, Français, Anglais")
    document.save(path)


if __name__ == "__main__":
    write_pdf(HERE / "sample_cv.pdf", PDF_LINES)
    write_docx(HERE / "sample_cv.docx")
    print("Wrote sample_cv.pdf and sample_cv.docx")
