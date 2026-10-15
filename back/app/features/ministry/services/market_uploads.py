"""Parsing and preview validation for normalized CSV/XLSX market datasets."""

import csv
import json
from datetime import date, datetime
from io import BytesIO, StringIO

from fastapi import HTTPException
from pydantic import ValidationError

from mahara_data.schemas.market import MarketIndicatorRecord

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_REJECTIONS = 5000


def _cell(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _records_csv(content: bytes) -> list[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=422, detail="CSV must use UTF-8 encoding") from error
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
        raise HTTPException(status_code=422, detail="CSV needs one row of unique column headers")
    result = []
    for row in reader:
        if not any(value not in (None, "") for value in row.values()):
            continue
        result.append({key: _cell(value) for key, value in row.items() if key is not None})
        if len(result) > MAX_ROWS:
            raise HTTPException(status_code=413, detail=f"Files are limited to {MAX_ROWS} data rows")
    return result


def _records_xlsx(content: bytes) -> list[dict]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers or any(not isinstance(value, str) or not value.strip() for value in headers):
            raise HTTPException(status_code=422, detail="Spreadsheet needs a header row with column names")
        names = [value.strip() for value in headers]
        if len(names) != len(set(names)):
            raise HTTPException(status_code=422, detail="Spreadsheet headers must be unique")
        result = []
        for row in rows:
            if not any(value is not None for value in row):
                continue
            result.append(dict(zip(names, (_cell(value) for value in row))))
            if len(result) > MAX_ROWS:
                raise HTTPException(status_code=413, detail=f"Files are limited to {MAX_ROWS} data rows")
        return result
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=422, detail="Could not read the XLSX workbook") from error


def parse_market_file(filename: str, content: bytes) -> tuple[list[dict], list[dict]]:
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail=f"Files are limited to {MAX_UPLOAD_BYTES} bytes")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        raw_rows = _records_csv(content)
    elif suffix == "xlsx":
        raw_rows = _records_xlsx(content)
    else:
        raise HTTPException(status_code=415, detail="Upload a .csv or .xlsx file")
    if len(raw_rows) > MAX_ROWS:
        raise HTTPException(status_code=413, detail=f"Files are limited to {MAX_ROWS} data rows")

    accepted: list[dict] = []
    rejected: list[dict] = []
    for index, row in enumerate(raw_rows):
        if isinstance(row.get("extra"), str):
            try:
                row["extra"] = json.loads(row["extra"] or "{}")
            except json.JSONDecodeError:
                rejected.append({"index": index, "reason": "extra must contain valid JSON"})
                continue
        row = {key: value for key, value in row.items() if value not in (None, "")}
        try:
            accepted.append(MarketIndicatorRecord.model_validate(row).model_dump(mode="json"))
        except ValidationError as error:
            reason = "; ".join(
                f"{'.'.join(str(part) for part in issue['loc'])}: {issue['msg']}"
                for issue in error.errors(include_input=False)
            )
            rejected.append({"index": index, "reason": reason[:1000]})
    return accepted, rejected[:MAX_REJECTIONS]
